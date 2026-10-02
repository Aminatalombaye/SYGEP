# -*- coding: utf-8 -*-
"""Moteur de mise en forme : HTML -> PDF (Chromium) et HTML -> Word (python-docx)."""
import os, re
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GREY = RGBColor(0x55, 0x5F, 0x6B)
COLORS = {'ok': RGBColor(0x1B, 0x7F, 0x3B), 'ko': RGBColor(0xC0, 0x1F, 0x2F), 'man': RGBColor(0x8A, 0x6D, 0x00),
          'prio-c': RGBColor(0xC0, 0x1F, 0x2F), 'prio-m': RGBColor(0xB4, 0x5F, 0x06), 'prio-l': GREY}


# ----------------------------------------------------------------- PDF
def html_to_pdf(html_path, pdf_path, landscape=False):
    import subprocess
    script = os.environ.get('TOPDF_JS', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'topdf.js'))
    subprocess.run(['node', script, os.path.abspath(html_path), os.path.abspath(pdf_path)], check=True, cwd=os.path.dirname(script))


# ----------------------------------------------------------------- DOCX helpers
def shade(cell_or_par, hex_fill):
    el = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, '_tc') else cell_or_par._p.get_or_add_pPr()
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hex_fill); el.append(s)


def par_border(par, side='left', color='0E7C86', sz=24, space=6):
    pPr = par._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); e = OxmlElement(f'w:{side}')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color)
    b.append(e); pPr.append(b)


def field(par, instr, placeholder=''):
    r = par.add_run(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'begin'); r._r.append(f)
    r = par.add_run(); i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = instr; r._r.append(i)
    r = par.add_run(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'separate'); r._r.append(f)
    par.add_run(placeholder)
    r = par.add_run(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'end'); r._r.append(f)


def set_cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)


def table_borders(table, color='C9D1DB'):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for k in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color); b.append(e)
    tblPr.append(b)


class Docx:
    def __init__(self, titre_pied, landscape=False):
        self.d = Document(); self.landscape = landscape
        s = self.d.sections[0]
        if landscape:
            s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        else:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(1.8 if landscape else 2.0); s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
        self.width_cm = (s.page_width - s.left_margin - s.right_margin) / 360000
        st = self.d.styles
        n = st['Normal']; n.font.name = 'Calibri'; n.font.size = Pt(10.5)
        n.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri'); n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.12
        for lvl, (sz, col, before, after) in {1: (22, NAVY, 0, 10), 2: (15, NAVY, 16, 6), 3: (12.5, TEAL, 12, 4), 4: (11, NAVY, 10, 3)}.items():
            h = st[f'Heading {lvl}']; h.font.name = 'Calibri'; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = col
            h.element.rPr.rFonts.set(qn('w:asciiTheme'), 'minorHAnsi') if False else None
            rf = h.element.rPr.rFonts
            for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), 'Calibri')
            for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after); h.paragraph_format.keep_with_next = True
        st['Heading 1'].paragraph_format.page_break_before = True
        # pied de page
        fp = s.footer.paragraphs[0]; fp.text = ''
        r = fp.add_run(titre_pied + '    |    Page '); r.font.size = Pt(8.5); r.font.color.rgb = GREY
        field(fp, 'PAGE', '1')
        for rr in fp.runs: rr.font.size = Pt(8.5); rr.font.color.rgb = GREY
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # demander à Word de mettre à jour les champs (sommaire) à l'ouverture
        up = OxmlElement('w:updateFields'); up.set(qn('w:val'), 'true'); self.d.settings.element.append(up)

    # ---- cover
    def cover(self, titre, sous_titre, organisme, lignes):
        d = self.d
        for _ in range(6): d.add_paragraph()
        p = d.add_paragraph(); r = p.add_run(titre); r.font.size = Pt(40); r.bold = True; r.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(2)
        p = d.add_paragraph(); par_border(p, 'top', '0E7C86', 36, 8); r = p.add_run(sous_titre); r.font.size = Pt(18); r.font.color.rgb = TEAL
        p.paragraph_format.space_before = Pt(8)
        p = d.add_paragraph(); r = p.add_run(organisme); r.font.size = Pt(12); r.font.color.rgb = GREY
        for _ in range(9): d.add_paragraph()
        for k, v in lignes:
            p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            a = p.add_run(k + '  '); a.bold = True; a.font.color.rgb = NAVY; a.font.size = Pt(10.5)
            b = p.add_run(v); b.font.size = Pt(10.5); b.font.color.rgb = GREY
        d.add_page_break()

    def toc(self, entries):
        """entries: [(niveau, titre)] affichés tels quels ; Word recalcule avec les numéros de page à l'ouverture."""
        d = self.d
        p = d.add_paragraph(); r = p.add_run('Sommaire'); r.font.size = Pt(22); r.bold = True; r.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(10)
        first = True
        for lvl, t in entries:
            p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2 if lvl > 1 else 3)
            p.paragraph_format.left_indent = Cm(0.7 * (lvl - 1))
            if first:
                r = p.add_run(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'begin'); r._r.append(f)
                r = p.add_run(); i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = 'TOC \\o "1-2" \\h \\z \\u'; r._r.append(i)
                r = p.add_run(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'separate'); r._r.append(f)
                first = False
            r = p.add_run(t); r.bold = lvl == 1; r.font.size = Pt(11 if lvl == 1 else 10)
            last = p
        r = last.add_run(); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'end'); r._r.append(f)
        self.d.add_page_break()

    # ---- inline
    def inline(self, par, node, fmt=None):
        fmt = dict(fmt or {})
        if isinstance(node, NavigableString):
            t = str(node)
            if not t: return
            r = par.add_run(t)
            self._fmt(r, fmt); return
        if not isinstance(node, Tag): return
        n = node.name
        if n == 'br': par.add_run().add_break(); return
        f = dict(fmt)
        if n in ('strong', 'b'): f['b'] = True
        if n in ('em', 'i'): f['i'] = True
        if n == 'code': f['code'] = True
        if n == 'a': f['link'] = True
        if n == 'span':
            for c in node.get('class', []):
                if c in COLORS: f['color'] = COLORS[c]; f['b'] = True
        for c in node.children: self.inline(par, c, f)

    def _fmt(self, r, f):
        if f.get('b'): r.bold = True
        if f.get('i'): r.italic = True
        if f.get('code'):
            r.font.name = 'Consolas'; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xA3, 0x1B, 0x5C)
            rPr = r._r.get_or_add_rPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'F1F3F5'); rPr.append(sh)
        if f.get('link'): r.font.color.rgb = RGBColor(0x1A, 0x5F, 0xB4); r.underline = True
        if f.get('color'): r.font.color.rgb = f['color']
        if f.get('size'): r.font.size = Pt(f['size'])

    # ---- blocks
    def blocks(self, container, quote=False, img_root='.'):
        for el in container.children:
            if isinstance(el, NavigableString):
                continue
            n = el.name
            if n in ('h1',): continue
            if n in ('h2', 'h3', 'h4', 'h5'):
                self.d.add_heading(el.get_text(' ', strip=True), level={'h2': 1, 'h3': 2, 'h4': 3, 'h5': 4}[n])
            elif n == 'p':
                p = self.d.add_paragraph()
                if quote: self._quote(p)
                for c in el.children: self.inline(p, c)
            elif n in ('ul', 'ol'):
                self.list(el, 0)
            elif n == 'table':
                self.table(el)
            elif n == 'pre':
                self.pre(el)
            elif n == 'blockquote':
                self.blocks(el, True, img_root)
            elif n == 'figure':
                self.figure(el, img_root)
            elif n == 'hr':
                pass
            elif n == 'div' and 'pagebreak' in el.get('class', []):
                self.d.add_page_break()
            elif n == 'div':
                self.blocks(el, quote, img_root)

    def _quote(self, p):
        p.paragraph_format.left_indent = Cm(0.5); par_border(p, 'left', '0E7C86', 24, 8); shade(p, 'EEF6F7')
        p.paragraph_format.space_before = Pt(2)

    def figure(self, fig, root):
        img = fig.find('img'); cap = fig.find('figcaption')
        src = os.path.join(root, img['src'])
        if os.path.exists(src):
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(src, width=Cm(min(15.5, self.width_cm - 0.5)))
        if cap:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(cap.get_text(strip=True)); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY

    def pre(self, el):
        lines = el.get_text().rstrip('\n').split('\n')
        for i, ln in enumerate(lines):
            p = self.d.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Cm(0.3)
            shade(p, 'F1F3F5')
            r = p.add_run(ln if ln else ' '); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
            r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Consolas')
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def list(self, el, depth):
        ordered = el.name == 'ol'
        for i, li in enumerate(el.find_all('li', recursive=False), 1):
            p = self.d.add_paragraph(style='List Bullet' if not ordered else None)
            p.paragraph_format.space_after = Pt(2)
            if ordered:
                p.paragraph_format.left_indent = Cm(0.9 + 0.7 * depth); p.paragraph_format.first_line_indent = Cm(-0.6)
                p.add_run(f'{i}.  ').bold = True
            elif depth:
                p.paragraph_format.left_indent = Cm(0.63 + 0.63 * depth)
            subs = []
            for c in li.children:
                if isinstance(c, Tag) and c.name in ('ul', 'ol'): subs.append(c)
                elif isinstance(c, Tag) and c.name == 'p':
                    if p.text.strip() and not ordered or (ordered and len(p.text) > len(f'{i}.  ')):
                        p = self.d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9 + 0.6 * depth)
                    for cc in c.children: self.inline(p, cc)
                else: self.inline(p, c)
            for s in subs: self.list(s, depth + 1)

    def table(self, el):
        rows = el.find_all('tr')
        ncol = max(len(r.find_all(['th', 'td'])) for r in rows)
        t = self.d.add_table(rows=0, cols=ncol); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        table_borders(t); set_cell_margins(t)
        head_cells = rows[0].find_all(['th', 'td'])
        ws = [float(c.get('data-w', 0)) for c in head_cells]
        if not all(ws):
            ws = [1.0] * ncol
        tot = sum(ws); cm = [self.width_cm * w / tot for w in ws]
        small = ncol >= 6 or self.landscape
        for ri, tr in enumerate(rows):
            cells = tr.find_all(['th', 'td'])
            row = t.add_row()
            if ri == 0:
                trPr = row._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
            cs = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cs)
            for ci in range(ncol):
                cell = row.cells[ci]; cell.width = Cm(cm[ci])
                if ci >= len(cells): continue
                src = cells[ci]; par = cell.paragraphs[0]; par.paragraph_format.space_after = Pt(0)
                fmt = {'size': 8.5 if small else 9.5}
                if ri == 0 and src.name == 'th':
                    shade(cell, '1F3A5F'); fmt.update(b=True, color=RGBColor(255, 255, 255))
                elif ri % 2 == 0:
                    shade(cell, 'F4F7FA')
                if src.get('class') and 'sec' in src.get('class', []):
                    shade(cell, 'DCE6F0')
                for c in src.children:
                    if isinstance(c, Tag) and c.name in ('ul', 'ol'):
                        for j, li in enumerate(c.find_all('li', recursive=False), 1):
                            if par.text: par = cell.add_paragraph(); par.paragraph_format.space_after = Pt(0)
                            self.inline(par, NavigableString(f'{j}. ' if c.name == 'ol' else '• '), fmt)
                            for cc in li.children: self.inline(par, cc, fmt)
                    elif isinstance(c, Tag) and c.name == 'p':
                        if par.text: par = cell.add_paragraph(); par.paragraph_format.space_after = Pt(0)
                        for cc in c.children: self.inline(par, cc, fmt)
                    else:
                        self.inline(par, c, fmt)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def save(self, path):
        PPR = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
               'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
               'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
               'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']
        TBL = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing',
               'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook']
        def reorder(root, tag, order):
            for el in root.iter(qn('w:' + tag)):
                kids = list(el)
                key = lambda k: order.index(k.tag.split('}')[1]) if k.tag.split('}')[1] in order else len(order)
                srt = sorted(kids, key=key)
                if srt != kids:
                    for k in kids: el.remove(k)
                    for k in srt: el.append(k)
        for part in [self.d.element.body, self.d.styles.element] + [sec.footer._element for sec in self.d.sections]:
            reorder(part, 'pPr', PPR); reorder(part, 'tblPr', TBL)
        self.d.save(path)
