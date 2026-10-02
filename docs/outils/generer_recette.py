# -*- coding: utf-8 -*-
"""Génère CAHIER-DE-RECETTE.md et CAHIER-DE-RECETTE.xlsx à partir de cas_de_test.py
et des résultats de la pré-recette automatique (pre_recette.json).
Usage : python3 docs/outils/generer_recette.py"""
import json, os, sys
sys.path.insert(0, os.path.dirname(__file__))
import cas_de_test as T
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

def cpt_txt(c):
    return f"{c} ({T.COMPTES[c][0]})" if c in T.COMPTES else c


ICI = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(ICI)
PRE = json.load(open(os.path.join(ICI, 'pre_recette.json'), encoding='utf-8'))
PRIO = {'C': 'Critique', 'M': 'Majeur', 'm': 'Mineur'}
ETAT = {'OK': 'OK', 'KO': 'KO', '—': 'Manuel'}

# ------------------------------------------------------------------ Markdown
md = ["# Cahier de recette SYGEP", "",
      "Document de validation fonctionnelle de l'application. Chaque cas est exécutable à la main par un testeur ; "
      "la colonne *Pré-recette* indique le résultat obtenu lors de l'exécution automatique (navigateur + requêtes HTTP) sur la base de démonstration.", "",
      "## Mode d'emploi", "",
      "1. Installer l'application avec les données de démonstration (`php artisan migrate --seed` puis `php artisan db:seed --class=DemoDataSeeder`).",
      "2. Pour chaque cas : se connecter avec le **compte** indiqué, vérifier les **préconditions**, suivre les **étapes**, comparer au **résultat attendu**.",
      "3. Renseigner le résultat (OK / KO / Bloqué / N/A), la date, le testeur et les observations (dans le fichier Excel).",
      "4. Priorités : **C** critique (bloquant pour la mise en production), **M** majeur, **m** mineur.",
      "5. Verdict : recette acceptée si 100 % des cas critiques sont OK et aucune anomalie bloquante ouverte.", "",
      "## Comptes de test (mot de passe de démonstration : `password`)", "",
      "| Code | E-mail | Rôle |", "|---|---|---|"]
for k, (e, r) in T.COMPTES.items():
    md.append(f"| {k} | {e} | {r} |")
tot = sum(len(m[3]) for m in T.MODULES)
auto_ok = sum(1 for v in PRE.values() if v[0] == 'OK')
auto_ko = sum(1 for v in PRE.values() if v[0] == 'KO')
md += ["", "## Synthèse de la pré-recette automatique", "",
       f"- Cas au cahier : **{tot}** ; vérifiés automatiquement : **{auto_ok + auto_ko}** (OK : {auto_ok}, KO : {auto_ko}) ; à faire à la main : **{tot - auto_ok - auto_ko}** (caméra, HTTPS, cron, sauvegarde, impression…).", "",
       "| Module | Cas | OK auto | KO auto | Manuel |", "|---|---|---|---|---|"]
for code, titre, obj, cas in T.MODULES:
    o = sum(1 for c in cas if PRE.get(c[0], ['', ''])[0] == 'OK')
    k = sum(1 for c in cas if PRE.get(c[0], ['', ''])[0] == 'KO')
    md.append(f"| {code} – {titre} | {len(cas)} | {o} | {k} | {len(cas) - o - k} |")
md += ["", "### Anomalies connues", ""]
for i, v in PRE.items():
    if v[0] == 'KO':
        md.append(f"- **{i}** : {v[1]}")
md += ["", "### Défauts trouvés et corrigés pendant la pré-recette", "",
       "- **Sécurité (IDOR)** : un utilisateur « local » pouvait ouvrir par URL directe une matière, une affectation ou un bon d'un autre service → désormais 404 (cas SEC-12).",
       "- Un utilisateur local ne pouvait pas transférer une matière à un agent de son service → corrigé (AFF-24/25).",
       "- Sortie de stock « par agent » : validation inatteignable → corrigée.",
       "- Alerte de rupture non émise quand le stock tombe à zéro → corrigée (STK-19).",
       "- Débordement horizontal sur mobile → corrigé.",
       "- Installation neuve : seeders non idempotents, rôle d'inscription libre mal ciblé → corrigés.", ""]
for code, titre, obj, cas in T.MODULES:
    md += [f"## {code} – {titre}", "", f"*Objectif : {obj}*", ""]
    for (i, test, cpt, pre, etapes, att, pr) in cas:
        r = PRE.get(i, ['—', ''])
        md += [f"### {i} · {test}", "",
               f"- **Compte** : {cpt_txt(cpt)} · **Priorité** : {PRIO[pr]} · **Pré-recette** : {ETAT.get(r[0], r[0])}",
               f"- **Préconditions** : {pre}", "- **Étapes** :"]
        md += [f"  {n}. {e}" for n, e in enumerate(etapes, 1)]
        md += [f"- **Résultat attendu** : {att}"]
        if r[1] and r[0] != '—':
            md.append(f"- **Constat pré-recette** : {r[1]}")
        md += ["- **Résultat** : ☐ OK ☐ KO ☐ Bloqué ☐ N/A — Date : ____ Testeur : ________", ""]
md += ["## Procès-verbal de recette", "",
       "| | |", "|---|---|", "| Application | SYGEP |", "| Version / commit | |", "| Date de recette | |",
       "| Cas exécutés / OK / KO | |", "| Anomalies bloquantes ouvertes | |", "| Décision | ☐ Acceptée ☐ Acceptée avec réserves ☐ Refusée |",
       "| Réserves | |", "", "Signatures : Maîtrise d'ouvrage ________  Maîtrise d'œuvre ________", ""]
open(os.path.join(DOCS, 'CAHIER-DE-RECETTE.md'), 'w', encoding='utf-8').write("\n".join(md))

# ------------------------------------------------------------------ Excel
wb = Workbook()
thin = Side(style='thin', color='BBBBBB'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill('solid', fgColor='1F3A5F'); HW = Font(bold=True, color='FFFFFF')
wrap = Alignment(wrap_text=True, vertical='top')

ws = wb.active; ws.title = "Mode d'emploi"
lignes = ["CAHIER DE RECETTE – SYGEP", "",
          "1. Une feuille par module. Une ligne par cas de test.",
          "2. Se connecter avec le compte indiqué (mot de passe de démonstration : password), suivre les étapes, comparer au résultat attendu.",
          "3. Choisir le résultat dans la liste (Non testé / OK / KO / Bloqué / N/A), puis renseigner date, testeur et observations.",
          "4. La feuille « Synthèse » se met à jour automatiquement. Les KO sont à reporter dans « Anomalies ».",
          "5. Priorités : Critique = bloquant ; Majeur ; Mineur. Acceptation : 100 % des cas critiques OK.",
          "6. La colonne « Pré-recette auto » donne le résultat de l'exécution automatique ; elle ne remplace pas votre test.", "",
          "COMPTES DE TEST"]
for n, l in enumerate(lignes, 1): ws.cell(n, 1, l)
ws['A1'].font = Font(bold=True, size=16); ws['A10'].font = Font(bold=True)
r0 = len(lignes) + 1
for k, (e, r) in T.COMPTES.items():
    ws.cell(r0, 1, f"{k}  –  {e}  –  {r}"); r0 += 1
ws.column_dimensions['A'].width = 130

heads = ["ID", "Test", "Compte", "Préconditions", "Étapes", "Résultat attendu", "Priorité", "Pré-recette auto",
         "Résultat", "Date", "Testeur", "Observations"]
widths = [9, 30, 22, 30, 52, 44, 10, 28, 12, 11, 14, 32]
sheets = []
for code, titre, obj, cas in T.MODULES:
    s = wb.create_sheet(code); sheets.append(s.title)
    s['A1'] = f"{code} – {titre}"; s['A1'].font = Font(bold=True, size=13)
    s['A2'] = obj
    for c, h in enumerate(heads, 1):
        x = s.cell(4, c, h); x.fill = HF; x.font = HW; x.border = B; x.alignment = Alignment(wrap_text=True, vertical='center')
        s.column_dimensions[x.column_letter].width = widths[c - 1]
    for r, (i, test, cpt, pre, etapes, att, pr) in enumerate(cas, 5):
        rr = PRE.get(i, ['—', ''])
        auto = {'OK': 'OK', 'KO': 'KO – ' + rr[1], '—': 'Manuel'}.get(rr[0], rr[0])
        vals = [i, test, f"{cpt_txt(cpt)}", pre, "\n".join(f"{n}. {e}" for n, e in enumerate(etapes, 1)),
                att, PRIO[pr], auto, "Non testé", None, None, None]
        for c, v in enumerate(vals, 1):
            x = s.cell(r, c, v); x.alignment = wrap; x.border = B
    last = 4 + len(cas)
    dv = DataValidation(type='list', formula1='"Non testé,OK,KO,Bloqué,N/A"', allow_blank=False)
    s.add_data_validation(dv); dv.add(f"I5:I{last}")
    for val, col in [('OK', 'C6EFCE'), ('KO', 'FFC7CE'), ('Bloqué', 'FFEB9C')]:
        s.conditional_formatting.add(f"I5:I{last}", CellIsRule(operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', bgColor=col, fgColor=col)))
    s.freeze_panes = 'C5'; s.auto_filter.ref = f"A4:L{last}"

sy = wb.create_sheet("Synthèse", 1)
for c, h in enumerate(["Module", "Total", "OK", "KO", "Bloqué", "N/A", "Non testé", "% OK", "Critiques non OK"], 1):
    x = sy.cell(1, c, h); x.fill = HF; x.font = HW
for r, ((code, titre, obj, cas), sn) in enumerate(zip(T.MODULES, sheets), 2):
    last = 4 + len(cas); rg = f"'{sn}'!$I$5:$I${last}"; pg = f"'{sn}'!$G$5:$G${last}"
    sy.cell(r, 1, f"{code} – {titre}"); sy.cell(r, 2, f"=COUNTA('{sn}'!$A$5:$A${last})")
    for c, v in zip(range(3, 8), ['OK', 'KO', 'Bloqué', 'N/A', 'Non testé']):
        sy.cell(r, c, f'=COUNTIF({rg},"{v}")')
    sy.cell(r, 8, f"=IF(B{r}=0,0,C{r}/B{r})").number_format = '0%'
    sy.cell(r, 9, f'=COUNTIFS({pg},"Critique",{rg},"<>OK",{rg},"<>N/A")')
n = len(T.MODULES) + 2
sy.cell(n, 1, "TOTAL").font = Font(bold=True)
for c in range(2, 8):
    L = sy.cell(1, c).column_letter; sy.cell(n, c, f"=SUM({L}2:{L}{n-1})").font = Font(bold=True)
sy.cell(n, 8, f"=IF(B{n}=0,0,C{n}/B{n})").number_format = '0%'
sy.cell(n, 9, f"=SUM(I2:I{n-1})")
sy.column_dimensions['A'].width = 52
for L in 'BCDEFGHI': sy.column_dimensions[L].width = 14

an = wb.create_sheet("Anomalies")
for c, h in enumerate(["N°", "Cas (ID)", "Description", "Gravité", "Statut", "Corrigée dans (commit)", "Date", "Commentaire"], 1):
    x = an.cell(1, c, h); x.fill = HF; x.font = HW
    an.column_dimensions[x.column_letter].width = [6, 12, 60, 12, 14, 22, 11, 40][c - 1]
r = 2
for i, v in PRE.items():
    if v[0] == 'KO':
        an.cell(r, 1, r - 1); an.cell(r, 2, i); an.cell(r, 3, v[1]); an.cell(r, 4, 'Majeur'); an.cell(r, 5, 'Ouverte'); r += 1
dv = DataValidation(type='list', formula1='"Ouverte,En cours,Corrigée,Fermée,Rejetée"'); an.add_data_validation(dv); dv.add("E2:E200")
dv2 = DataValidation(type='list', formula1='"Bloquante,Majeur,Mineur"'); an.add_data_validation(dv2); dv2.add("D2:D200")

pv = wb.create_sheet("PV de recette")
for r, (a, b) in enumerate([("Application", "SYGEP"), ("Version / commit", ""), ("Date de recette", ""), ("Testeurs", ""),
                            ("Cas exécutés", f"='Synthèse'!B{n}-Synthèse!G{n}"), ("Cas OK", f"='Synthèse'!C{n}"), ("Cas KO", f"='Synthèse'!D{n}"),
                            ("Critiques non OK", f"='Synthèse'!I{n}"), ("Décision (Acceptée / Avec réserves / Refusée)", ""),
                            ("Réserves", ""), ("Signature maîtrise d'ouvrage", ""), ("Signature maîtrise d'œuvre", "")], 1):
    pv.cell(r, 1, a).font = Font(bold=True); pv.cell(r, 2, b)
pv.column_dimensions['A'].width = 46; pv.column_dimensions['B'].width = 60
wb.save(os.path.join(DOCS, 'CAHIER-DE-RECETTE.xlsx'))
print("OK", tot, "cas")
